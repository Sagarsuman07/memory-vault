import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from PIL import Image
from types import SimpleNamespace

from ingestion.processing import clean_and_normalize_text, process_content
from ingestion.processing_models import ExtractionError
from ingestion.document_processor import extract_document_text, extract_pdf_content
from ingestion.image_processor import extract_image_content
from ingestion.audio_processor import extract_audio_content


class TestTextCleaning(unittest.TestCase):
    def test_normalizes_line_endings_unicode_and_excess_blank_lines(self):
        self.assertEqual(clean_and_normalize_text("A\r\n\r\n\r\nB\x00\u00a0C "), "A\n\nB C")

    def test_preserves_semantic_identifiers_negations_and_amounts(self):
        text = "Flight AB-123 is NOT cancelled. Price: ₹8,500. Date: 18/10/2026."
        self.assertEqual(clean_and_normalize_text(text), text)


class TestDocumentProcessing(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_txt_extracts_clean_text_and_encoding_metadata(self):
        path = self.root / "notes.txt"
        path.write_bytes(b"\xef\xbb\xbfHello\r\n\r\nWorld  ")
        result = process_content(path)
        self.assertEqual(result.text, "Hello\n\nWorld")
        self.assertEqual(result.metadata["document_type"], "text")
        self.assertEqual(result.metadata["encoding"], "utf-8-sig")
        self.assertEqual(extract_document_text(str(path)), result.text)

    def test_docx_preserves_headings_and_table_rows(self):
        path = self.root / "notes.docx"
        doc = Document()
        doc.add_heading("Travel Details", level=1)
        doc.add_paragraph("Flight AB123 is NOT cancelled.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Price"
        table.cell(1, 1).text = "₹8,500"
        doc.save(path)
        result = process_content(path)
        self.assertIn("## Travel Details", result.text)
        self.assertIn("Field | Value", result.text)
        self.assertIn("Price | ₹8,500", result.text)
        self.assertEqual(result.metadata["table_count"], 1)
        self.assertIn("Travel Details", result.metadata["headings"])

    def test_pdf_preserves_page_boundaries_and_metadata(self):
        path = self.root / "notes.pdf"
        doc = fitz.open()
        p1 = doc.new_page()
        p1.insert_text((72, 72), "Flight AB123 departs at 10:30 AM.")
        p2 = doc.new_page()
        p2.insert_text((72, 72), "Page two contains Price: 8500.")
        doc.set_metadata({"title": "Travel Test", "author": "Test Author"})
        doc.save(path)
        doc.close()
        result = process_content(path)
        self.assertIn("[Page 1]", result.text)
        self.assertIn("[Page 2]", result.text)
        self.assertEqual(result.metadata["page_count"], 2)
        self.assertEqual(result.metadata["title"], "Travel Test")
        self.assertEqual(result.metadata["pages_with_text"], 2)

    def test_scanned_pdf_without_ocr_fails_explicitly(self):
        path = self.root / "scanned.pdf"
        doc = fitz.open()
        doc.new_page()
        doc.save(path)
        doc.close()
        with patch("ingestion.document_processor.settings") as mock_settings:
            mock_settings.PDF_OCR_ENABLED = False
            with self.assertRaisesRegex(ExtractionError, "OCR is disabled"):
                extract_pdf_content(str(path))

    def test_pdf_ocr_fallback_is_used_when_enabled(self):
        path = self.root / "scanned.pdf"
        doc = fitz.open()
        doc.new_page()
        doc.save(path)
        doc.close()
        with patch("ingestion.document_processor.settings") as mock_settings, \
             patch("ingestion.document_processor._ocr_pdf_page", return_value="Recognized ticket AB123") as mock_ocr:
            mock_settings.PDF_OCR_ENABLED = True
            mock_settings.PDF_OCR_DPI = 144
            mock_settings.PDF_OCR_MIN_PAGE_TEXT_CHARS = 20
            result = extract_pdf_content(str(path))
        self.assertIn("Recognized ticket AB123", result.text)
        self.assertEqual(result.status, "succeeded")
        self.assertTrue(mock_ocr.called)
        self.assertEqual(result.metadata["extraction_method"], "pymupdf_text+ollama_vision_ocr")

    def test_empty_txt_is_rejected(self):
        path = self.root / "empty.txt"
        path.write_text("  ", encoding="utf-8")
        with self.assertRaises(ExtractionError):
            process_content(path)

    def test_image_extraction_returns_vision_metadata(self):
        path = self.root / "test.png"
        Image.new("RGB", (20, 10), "white").save(path)
        fake_response = SimpleNamespace(message=SimpleNamespace(content="Receipt total: ₹500"))
        fake_client = SimpleNamespace(chat=lambda **kwargs: fake_response)
        with patch("ingestion.image_processor.get_ollama_client", return_value=fake_client):
            result = extract_image_content(str(path))
        self.assertIn("₹500", result.text)
        self.assertEqual(result.metadata["image_width"], 20)
        self.assertEqual(result.metadata["image_height"], 10)
        self.assertEqual(result.metadata["extraction_method"], "ollama_vision")

    def test_audio_extraction_returns_transcription_metadata(self):
        path = self.root / "voice.wav"
        path.write_bytes(b"test audio bytes")
        fake_transcription = SimpleNamespace(text="Remember the meeting at 3 PM.")
        fake_client = SimpleNamespace(audio=SimpleNamespace(transcriptions=SimpleNamespace(
            create=lambda **kwargs: fake_transcription)))
        with patch("ingestion.audio_processor.get_groq_client", return_value=fake_client):
            result = extract_audio_content(str(path))
        self.assertIn("3 PM", result.text)
        self.assertEqual(result.metadata["document_type"], "audio")
        self.assertEqual(result.metadata["transcription_model"], "whisper-large-v3-turbo")


if __name__ == "__main__":
    unittest.main()
